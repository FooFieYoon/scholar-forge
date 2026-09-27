# -*- coding: utf-8 -*-
"""OOXML 收尾规范化：pPr/rPr/tcPr/trPr/tblPr 子元素按 ECMA-376 顺序重排 ＋ 补齐缺 eastAsia。

用法：
    python normalize_ooxml.py <文档.docx> [--check-only]

为什么必须做：手写 OOXML 时若目标元素「已存在」，顺序插入函数不会重定位它，
原稿自带的不合规顺序（常见于工具生成的 docx）会残留。本脚本全量重排一遍即清零。
Word 对 pPr/rPr 子元素顺序有一定容忍度，但清零可消除"文档需修复"提示的风险。
"""
import sys
import zipfile
import argparse

import docx
from docx.oxml.ns import qn
from lxml import etree

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'

PPR_ORDER = ['w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr',
             'w:widowControl', 'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd', 'w:tabs',
             'w:suppressAutoHyphens', 'w:kinsoku', 'w:wordWrap', 'w:overflowPunct',
             'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd',
             'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
             'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
             'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr',
             'w:pPrChange']
RPR_ORDER = ['w:rStyle', 'w:rFonts', 'w:b', 'w:bCs', 'w:i', 'w:iCs', 'w:caps', 'w:smallCaps',
             'w:strike', 'w:dstrike', 'w:outline', 'w:shadow', 'w:emboss', 'w:imprint',
             'w:noProof', 'w:snapToGrid', 'w:vanish', 'w:webHidden', 'w:color', 'w:spacing',
             'w:w', 'w:kern', 'w:position', 'w:sz', 'w:szCs', 'w:highlight', 'w:u', 'w:effect',
             'w:bdr', 'w:shd', 'w:fitText', 'w:vertAlign', 'w:rtl', 'w:cs', 'w:em', 'w:lang',
             'w:eastAsianLayout', 'w:specVanish', 'w:oMath']
TCPR_ORDER = ['w:cnfStyle', 'w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge', 'w:tcBorders',
              'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection', 'w:tcFitText', 'w:vAlign',
              'w:hideMark', 'w:headers', 'w:cellIns', 'w:cellDel', 'w:cellMerge', 'w:tcPrChange']
TRPR_ORDER = ['w:cnfStyle', 'w:divId', 'w:gridBefore', 'w:gridAfter', 'w:wBefore', 'w:wAfter',
              'w:cantSplit', 'w:trHeight', 'w:tblHeader', 'w:tblCellSpacing', 'w:jc', 'w:hidden',
              'w:ins', 'w:del', 'w:trPrChange']
TBLPR_ORDER = ['w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual', 'w:tblStyleRowBandSize',
               'w:tblStyleColBandSize', 'w:tblW', 'w:jc', 'w:tblCellSpacing', 'w:tblInd',
               'w:tblBorders', 'w:shd', 'w:tblLayout', 'w:tblCellMar', 'w:tblLook',
               'w:tblCaption', 'w:tblDescription']

TH = {'pPr': PPR_ORDER, 'rPr': RPR_ORDER, 'tcPr': TCPR_ORDER, 'trPr': TRPR_ORDER,
      'tblPr': TBLPR_ORDER}


def key_of(tag):
    return tag.replace(W, 'w:')


def violations(root):
    n = 0
    for name, order in TH.items():
        for el in root.iter(qn('w:' + name)):
            ks = [key_of(c.tag) for c in el if isinstance(c.tag, str)]
            idx = [order.index(k) for k in ks if k in order]
            if idx != sorted(idx):
                n += 1
    return n


def normalize(root):
    """稳定重排：已知元素按 schema 顺序，未知元素保持原相对次序置于末尾。"""
    fixed = 0
    for name, order in TH.items():
        for el in root.iter(qn('w:' + name)):
            kids = [c for c in el if isinstance(c.tag, str)]
            ks = [key_of(c.tag) for c in kids]
            idx = [order.index(k) for k in ks if k in order]
            if idx == sorted(idx):
                continue
            unknown = [c for c, k in zip(kids, ks) if k not in order]
            known = sorted([c for c, k in zip(kids, ks) if k in order],
                           key=lambda c: order.index(key_of(c.tag)))
            for c in kids:
                el.remove(c)
            for c in known + unknown:
                el.append(c)
            fixed += 1
    return fixed


def _hf_parts(doc):
    """只取已定义的页眉/页脚部件——访问 hf.paragraphs 会为"链接到前节"的节凭空创建空部件。"""
    out = []
    for s in doc.sections:
        for hf in (s.footer, s.header):
            try:
                if hf.is_linked_to_previous:
                    continue
                for p in hf.paragraphs:
                    out.append(p._p.getroottree().getroot())
            except Exception:
                pass
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('docx_path')
    ap.add_argument('--check-only', action='store_true')
    ap.add_argument('--east-asia', default='宋体')
    args = ap.parse_args()

    z = zipfile.ZipFile(args.docx_path)
    bad = z.testzip()
    print('ZIP testzip() ->', bad if bad else 'None（完整）')
    for n in z.namelist():
        if n.endswith('.xml') or n.endswith('.rels'):
            try:
                etree.fromstring(z.read(n))
            except Exception as ex:
                print('  !! %s 解析失败: %s' % (n, ex))
    print('全 XML 解析通过')

    doc = docx.Document(args.docx_path)
    parts = [doc.element.body, doc.styles.element] + _hf_parts(doc)
    print('修正前违规 %d 处' % sum(violations(p) for p in parts))
    if args.check_only:
        return
    tot = sum(normalize(p) for p in parts)
    print('已重排 %d 处' % tot)

    miss = 0
    for r in doc.element.body.iter(qn('w:r')):
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            continue
        rf = rpr.find(qn('w:rFonts'))
        if rf is not None and rf.get(qn('w:eastAsia')) is None:
            rf.set(qn('w:eastAsia'), args.east_asia)
            miss += 1
    print('补齐缺 eastAsia 的 run %d 个' % miss)
    doc.save(args.docx_path)
    print('修正后违规 %d 处' % sum(violations(p) for p in parts))
    print('已保存:', args.docx_path)


if __name__ == '__main__':
    main()
