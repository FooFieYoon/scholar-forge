# -*- coding: utf-8 -*-
"""把可编辑版 pptx 近似渲染成图片，用于与原视频帧对照 + OCR 校验。"""
import os, io, re, math
from pptx import Presentation
from pptx.util import Emu
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from rapidocr_onnxruntime import RapidOCR

W = r"E:/User/Documents/Dev_Projects/WorkBuddy/PPT2/work"
PPTX = r"E:/User/Documents/Dev_Projects/WorkBuddy/PPT2/光的反射-说课（可编辑版）.pptx"
PREV2 = os.path.join(W, "preview2")
os.makedirs(PREV2, exist_ok=True)
W_PX, H_PX = 1106, 648
PX_PER_IN = 1106 / 13.3333
EMUX = 914400
SIM = "C:/Windows/Fonts/simsun.ttc"
HEI = "C:/Windows/Fonts/simhei.ttf"


def emu2px(v):
    return v / EMUX * PX_PER_IN


def alpha_of(shape):
    try:
        e = shape.fill.fore_color._xFill.find(qn("a:srgbClr")).find(qn("a:alpha"))
        return int(e.get("val")) / 100000.0 if e is not None else 1.0
    except Exception:
        return 1.0


def rgb_of(shape):
    try:
        c = shape.fill.fore_color.rgb
        return (c[0], c[1], c[2])
    except Exception:
        return None


def line_of(shape):
    try:
        c = shape.line.color.rgb
        w = shape.line.width.pt if shape.line.width else 1
        return (c[0], c[1], c[2]), max(1, int(round(w * 1.2)))
    except Exception:
        return None


def para_lines(p):
    lines, cur = [], []
    for ch in p._p:
        tag = ch.tag.split("}")[1]
        if tag == "br":
            lines.append("".join(cur)); cur = []
        elif tag == "r":
            cur.append("".join(n.text or "" for n in ch.iter() if n.tag.endswith("}t")))
    lines.append("".join(cur))
    return lines


def draw_text(dr, tf, x, y, w, h, default_color=(0, 0, 0)):
    anchor = tf.vertical_anchor
    paras = []
    for p in tf.paragraphs:
        if not p.runs:
            continue
        size = p.runs[0].font.size.pt if p.runs[0].font.size else 18
        bold = bool(p.runs[0].font.bold)
        col = p.runs[0].font.color.rgb if p.runs[0].font.color and p.runs[0].font.color.rgb else default_color
        col = (col[0], col[1], col[2])
        ls = p.line_spacing.pt if hasattr(p.line_spacing, "pt") else size * 1.2
        align = p.alignment
        paras.append((para_lines(p), size, bold, col, ls, align))
    if not paras:
        return
    total = sum(len(ls_) * (s * 1.2 if lsp is None else lsp * 1.2)
                for ls_, s, b, c, lsp, a in paras)
    cy = y
    if str(anchor) == "MSO_ANCHOR.MIDDLE" or (hasattr(anchor, "real") and anchor == 2):
        cy = y + (h - total) / 2
    for ls_, size, bold, col, lsp, align in paras:
        font = ImageFont.truetype(HEI if bold else SIM, max(6, int(round(size * 1.2))))
        step = (lsp or size * 1.2) * 1.2
        for ln in ls_:
            if not ln:
                cy += step
                continue
            tw = dr.textlength(ln, font=font)
            tx = x
            if align is not None and "CENTER" in str(align):
                tx = x + (w - tw) / 2
            dr.text((tx, cy), ln, font=font, fill=col)
            cy += step


prs = Presentation(PPTX)
for i, slide in enumerate(prs.slides, 1):
    img = Image.new("RGB", (W_PX, H_PX), (245, 245, 243))
    overlay = Image.new("RGBA", (W_PX, H_PX), (0, 0, 0, 0))
    dr = ImageDraw.Draw(img)
    texts = []
    for sh in sorted(slide.shapes, key=lambda s: s.shape_id):
        x, y = emu2px(sh.left), emu2px(sh.top)
        w, h = emu2px(sh.width), emu2px(sh.height)
        if sh.shape_type == 13:        # PICTURE
            try:
                p = Image.open(io.BytesIO(sh.image.blob)).convert("RGB").resize(
                    (max(1, int(w)), max(1, int(h))))
                img.paste(p, (int(x), int(y)))
            except Exception as e:
                print("pic err", e)
            continue
        # autoshape
        try:
            prst = sh._element.find(qn("p:spPr")).find(qn("a:prstGeom")).get("prst")
        except Exception:
            prst = "rect"
        col = rgb_of(sh)
        a = alpha_of(sh)
        linfo = line_of(sh)
        box = [x, y, x + w, y + h]
        od = ImageDraw.Draw(overlay)
        if prst in ("roundRect", "rect"):
            if col:
                od.rectangle(box, fill=col + (int(255 * a),))
            if linfo:
                dr.rectangle(box, outline=linfo[0], width=linfo[1])
        elif prst == "ellipse":
            if col:
                od.ellipse(box, fill=col + (int(255 * a),))
            if linfo:
                dr.ellipse(box, outline=linfo[0], width=linfo[1])
        elif prst == "diamond":
            pts = [(x + w / 2, y), (x + w, y + h / 2), (x + w / 2, y + h), (x, y + h / 2)]
            if col:
                od.polygon(pts, fill=col + (int(255 * a),))
        if sh.has_text_frame and sh.text_frame.text.strip():
            texts.append((sh.text_frame, x, y, w, h))
    img = Image.alpha_composite(img.convert("RGBA"), overlay).convert("RGB")
    dr = ImageDraw.Draw(img)
    for tf, x, y, w, h in texts:            # text always on top
        draw_text(dr, tf, x, y, w, h)
    img.save(os.path.join(PREV2, f"p{i:02d}.png"))
print("rendered", len(prs.slides.__iter__.__self__._sldIdLst), "slides ->", PREV2)

ocr = RapidOCR()
src = None
for i in range(1, 18):
    res, _ = ocr(os.path.join(PREV2, f"p{i:02d}.png"))
    txt = " | ".join(x[1] for x in res) if res else ""
    print(f"p{i:02d}: {txt[:170]}")
