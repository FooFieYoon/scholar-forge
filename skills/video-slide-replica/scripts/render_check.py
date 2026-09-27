import json, os, re
from pptx import Presentation
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
from rapidocr_onnxruntime import RapidOCR

W = r"E:/User/Documents/Dev_Projects/WorkBuddy/PPT2/work"
PPTX = r"E:/User/Documents/Dev_Projects/WorkBuddy/PPT2/光的反射-说课（复刻版）.pptx"
PREV = os.path.join(W, "preview")
os.makedirs(PREV, exist_ok=True)
PX_PER_IN = 1106 / 13.3333
EMU_PER_IN = 914400
W_PX, H_PX = 1106, 648
FONTS = {"宋体": r"C:\Windows\Fonts\simsun.ttc", "simsun": r"C:\Windows\Fonts\simsun.ttc"}


def load_font(size, bold=False):
    path = r"C:\Windows\Fonts\simhei.ttf" if bold else r"C:\Windows\Fonts\simsun.ttc"
    return ImageFont.truetype(path, max(6, int(round(size))))


def para_lines(p):
    lines = [[]]
    for child in p._p:
        tag = child.tag.split("}")[1]
        if tag == "br":
            lines.append([])
        elif tag == "r":
            txt = "".join(n.text or "" for n in child.iter() if n.tag.endswith("}t"))
            lines[-1].append(txt)
    return ["".join(l) for l in lines]


prs = Presentation(PPTX)
origin = json.load(open(os.path.join(W, "pages_text.json"), encoding="utf-8"))
ocr = RapidOCR()

summary = []
for i, slide in enumerate(prs.slides, 1):
    bg = Image.open(os.path.join(W, "pages_bg", f"p{i:02d}.png")).convert("RGB")
    bg = bg.resize((W_PX, H_PX))
    dr = ImageDraw.Draw(bg)
    for sh in slide.shapes:
        if not sh.has_text_frame or not sh.text_frame.text.strip():
            continue
        x = sh.left / EMU_PER_IN * PX_PER_IN
        y = sh.top / EMU_PER_IN * PX_PER_IN
        bodyPr = sh.text_frame._txBody.find(qn("a:bodyPr"))
        vert = bodyPr is not None and bodyPr.get("vert") == "eaVert"
        for p in sh.text_frame.paragraphs:
            if not p.runs:
                continue
            size_pt = p.runs[0].font.size.pt if p.runs[0].font.size else 18
            bold = bool(p.runs[0].font.bold)
            rgb = p.runs[0].font.color.rgb
            col = (rgb[0], rgb[1], rgb[2]) if rgb else (0, 0, 0)
            font = load_font(size_pt * PX_PER_IN / 72.0 * 72 / 72 * 1, bold)
            font = load_font(size_pt * (PX_PER_IN / 72.0), bold)
            step = size_pt * (PX_PER_IN / 72.0) * 1.0
            for k, line in enumerate(para_lines(p)):
                yy = y + k * (p.line_spacing.pt * (PX_PER_IN / 72.0) if hasattr(p.line_spacing, "pt") else step)
                if vert:
                    for c_i, ch in enumerate(line):
                        dr.text((x, yy + c_i * step * 1.02), ch, font=font, fill=col)
                else:
                    dr.text((x, yy), line, font=font, fill=col)
    out = os.path.join(PREV, f"p{i:02d}.png")
    bg.save(out)
    res, _ = ocr(out)
    got = [(x[1], [int(min(p[0] for p in x[0])), int(min(p[1] for p in x[0]))]) for x in res] if res else []
    src = [t for t in origin[i - 1]["texts"]]
    diffs = []
    for t in src:
        key = re.sub(r"\s+", "", t["t"]).replace("人射", "入射").replace("反财", "反射").replace("反时", "反射")
        best, bd = None, 1e9
        for gt, gp in got:
            gk = re.sub(r"\s+", "", gt)
            if not gk:
                continue
            common = len(set(key) & set(gk)) / max(1, len(set(key) | set(gk)))
            if key in gk or gk in key:
                common = 1.0
            if common > 0.75:
                d = abs(gp[0] - t["x"]) + abs(gp[1] - t["y"])
                if d < bd:
                    best, bd = (gt, gp), d
        diffs.append(dict(text=t["t"][:22], src=[t["x"], t["y"]],
                          got=best[1] if best else None, d=round(bd) if best else None))
    ok = [d for d in diffs if d["d"] is not None]
    miss = [d for d in diffs if d["d"] is None]
    summary.append(dict(page=i, matched=len(ok), missing=len(miss),
                        mean_dev=round(sum(d["d"] for d in ok) / len(ok), 1) if ok else None,
                        max_dev=max((d["d"] for d in ok), default=None),
                        missing_texts=[m["text"] for m in miss]))
    print(f"p{i:02d} matched={len(ok):2d}/{len(diffs):2d} meanDev="
          f"{summary[-1]['mean_dev']} maxDev={summary[-1]['max_dev']} "
          f"missing={summary[-1]['missing_texts']}")
json.dump(summary, open(os.path.join(W, "render_check.json"), "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
