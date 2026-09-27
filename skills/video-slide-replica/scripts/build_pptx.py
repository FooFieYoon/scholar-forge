import json, os, re, copy
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn

W = r"E:/User/Documents/Dev_Projects/WorkBuddy/PPT2/work"
OUT = r"E:/User/Documents/Dev_Projects/WorkBuddy/PPT2/光的反射-说课（复刻版）.pptx"
BG = os.path.join(W, "pages_bg")
data = json.load(open(os.path.join(W, "pages_text.json"), encoding="utf-8"))

SLIDE_W_IN, SLIDE_H_IN = 13.3333, 7.8125
PX2IN = SLIDE_W_IN / 1106.0
PX2PT = PX2IN * 72.0
FONT = "宋体"

TEXT_FIX = {
    "2.继光的直线倒摊后对光现象": "2.继光的直线传播后对光现象",
    "2.继光的直线倒摊后对光现": "2.继光的直线传播后对光现象",
}
DROP = {(4, "维能力有限"), (4, "时抽象概念的理解存"), (4, "维能力有限")}

prs = Presentation()
prs.slide_width = Inches(SLIDE_W_IN)
prs.slide_height = Inches(SLIDE_H_IN)
blank = prs.slide_layouts[6]


def set_font(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def add_text(slide, lines, x, y, w, size_pt, align=PP_ALIGN.LEFT, line_pt=None,
             vert=False, max_h=None):
    pad = 3
    h = max_h if max_h else 40
    tb = slide.shapes.add_textbox(Inches((x - pad) * PX2IN), Inches((y - pad) * PX2IN),
                                  Inches((w + 6) * PX2IN), Inches((h + 2 * pad) * PX2IN))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = Pt(line_pt) if line_pt else 1.0
    for i, (txt, rgb, bold) in enumerate(lines):
        if i:
            p._p.append(p._p.makeelement(qn("a:br"), {}))
        r = p.add_run()
        r.text = txt
        r.font.size = Pt(size_pt)
        r.font.bold = bold
        r.font.color.rgb = RGBColor(*rgb)
        set_font(r)
    if vert:
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        bodyPr.set("vert", "eaVert")
    return tb


def iou_min(a, b):
    ax0, ay0, ax1, ay1 = a["x"], a["y"], a["x"] + a["w"], a["y"] + a["h"]
    bx0, by0, bx1, by1 = b["x"], b["y"], b["x"] + b["w"], b["y"] + b["h"]
    iw = max(0, min(ax1, bx1) - max(ax0, bx0))
    ih = max(0, min(ay1, by1) - max(ay0, by0))
    im = iw * ih
    if im == 0:
        return 0.0
    return im / min(max(1, a["w"] * a["h"]), max(1, b["w"] * b["h"]))


def clean(s):
    return re.sub(r"\s+", " ", s.replace("#", "")).strip()


def prep(pidx, texts):
    out = []
    for t in texts:
        s = TEXT_FIX.get(clean(t["t"]), clean(t["t"]))
        if not s or (pidx, s) in DROP:
            continue
        d = dict(t=s, **{k: v for k, v in t.items() if k != "t"})
        out.append(d)
    return out


def dedup(texts):
    texts = sorted(texts, key=lambda a: -(a["w"] * a["h"]))
    keep = []
    for t in texts:
        dup = False
        for k in keep:
            if iou_min(t, k) > 0.45:
                dup = True
                if len(t["t"]) > len(k["t"]) + 1 and t["t"] not in k["t"]:
                    k.update(t)
                break
        if not dup:
            keep.append(t)
    return keep


def group_paragraphs(texts):
    items = sorted(texts, key=lambda a: (a["y"], a["x"]))
    used = [False] * len(items)
    groups = []
    for i, a in enumerate(items):
        if used[i]:
            continue
        grp = [a]
        used[i] = True
        changed = True
        while changed:
            changed = False
            last = grp[-1]
            for j, b in enumerate(items):
                if used[j]:
                    continue
                if abs(b["x"] - grp[0]["x"]) < 14 and \
                   abs(b["h"] - last["h"]) < max(6, last["h"] * 0.35) and \
                   last["y"] < b["y"] <= last["y"] + last["h"] * 1.6 + 8:
                    grp.append(b)
                    used[j] = True
                    changed = True
                    break
        groups.append(sorted(grp, key=lambda a: a["y"]))
    return groups


OVERRIDE = {
    1: [dict(t="人教版  八年级物理  第四章第二节", x=39, y=87, w=563, h=30, ink=[107, 120, 83]),
        dict(t="《光的反射》说课", x=184, y=278, w=680, h=91, ink=[161, 76, 39])],
}
EXTRA = {16: [dict(t="3.", x=381, y=415, w=45, h=38, ink=[253, 253, 251])]}

report = {}
for page in data:
    idx = page["index"]
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(os.path.join(BG, f"p{idx:02d}.png"), 0, 0,
                             width=prs.slide_width, height=prs.slide_height)
    if idx in OVERRIDE:
        texts = copy.deepcopy(OVERRIDE[idx])
    else:
        texts = dedup(prep(idx, page["texts"]))
        if idx == 16:
            out = []
            for t in texts:
                if t["t"].startswith("2.") and t["h"] > 50:
                    out.append(dict(t="2.", x=t["x"], y=t["y"], w=45, h=38, ink=t["ink"]))
                else:
                    out.append(t)
            texts = out + copy.deepcopy(EXTRA[16])

    lines_report = []
    for grp in group_paragraphs(texts):
        first = grp[0]
        import statistics
        size = statistics.median([g["h"] for g in grp]) * PX2PT * 0.90
        lines = []
        for g in grp:
            col = [max(0, min(255, int(c))) for c in g["ink"]]
            bold = (first["x"] < 110 and first["y"] < 90)
            lines.append((g["t"], col, bold))
        wpx = max(g["w"] for g in grp)
        line_pt = None
        if len(grp) > 1:
            line_pt = (grp[-1]["y"] - first["y"]) / (len(grp) - 1) * PX2PT
        vert = first["h"] > wpx * 1.8
        align = PP_ALIGN.LEFT
        if idx == 1 and first["y"] > 200:
            size = first["h"] * PX2PT * 0.92
            lines = [(t, c, True) for t, c, _ in lines]
            align = PP_ALIGN.CENTER
        add_text(slide, lines, first["x"], first["y"], wpx, size,
                 align=align, line_pt=line_pt, vert=vert, max_h=first["h"])
        lines_report.append(dict(lines=[l[0] for l in lines], x=first["x"], y=first["y"],
                                 size=round(size, 1), vert=vert))
    report[idx] = lines_report

prs.save(OUT)
json.dump(report, open(os.path.join(W, "build_report.json"), "w", encoding="utf-8"),
          ensure_ascii=False, indent=1)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")
for k, v in report.items():
    print(f"p{k}: {len(v)} text objects; sizes", sorted({round(x['size']) for x in v}))
