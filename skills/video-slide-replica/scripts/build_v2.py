# -*- coding: utf-8 -*-
"""16:9 可编辑版复刻：所有元素用 pptx 原生形状/文本框重绘，照片为独立图片对象。"""
import os, json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
try:
    from pptx.oxml.xmlchemy import OxmlElement
except ImportError:                       # older / newer layout
    from pptx.oxml import parse_xml
    from pptx.oxml.ns import nsdecls

    def OxmlElement(tag, nsmap=None):
        return parse_xml(f"<{tag} {nsdecls('a')}/>")
import cv2

W = r"E:/User/Documents/Dev_Projects/WorkBuddy/PPT2/work"
OUT = r"E:/User/Documents/Dev_Projects/WorkBuddy/PPT2/光的反射-说课（可编辑版）.pptx"
ASSETS = os.path.join(W, "assets")
os.makedirs(ASSETS, exist_ok=True)

# ---- geometry: source slide bitmap is 1106 x 648, target slide 13.333 x 7.5 in (16:9)
SW, SH = 13.3333, 7.5
XI = SW / 1106.0
YI = SH / 648.0
PT = YI * 72.0 * 0.94          # px -> pt  (keeps the visual weight of the original)
FONT = "宋体"

# palette sampled from the video
C_BG      = "F5F5F3"
C_GREEN   = "6D8250"   # 墨绿卡片
C_GREENL  = "718454"   # 墨绿装饰/描边
C_TITLE   = "6F8250"   # 页眉墨绿
C_ORANGE  = "A14C27"   # 封面主标题/强调橙红
C_OCHRE   = "7A310E"   # 目录/小标题赭色
C_TEXT    = "1A1A1A"   # 正文黑
C_RED     = "96050A"   # 高亮红
C_WHITE   = "FFFFFF"
C_CYAN    = "C2EBEA"   # 环节四浅青块
C_BOARD   = "334C38"   # 板书深墨绿

prs = Presentation()
prs.slide_width = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


def hx(s):
    return RGBColor(int(s[0:2], 16), int(s[2:4], 16), int(s[4:6], 16))


def alpha(shape, pct):
    """fill transparency, eg. 0.4"""
    try:
        clr = shape.fill.fore_color._xFill.find(qn("a:srgbClr"))
    except Exception:
        return
    if clr is None:
        return
    a = OxmlElement("a:alpha")
    a.set("val", str(int((1 - pct) * 100000)))
    clr.append(a)


def rect(slide, x, y, w, h, fill=None, round_=0, line=None, lw=1.0, alpha_pct=None):
    shp = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE,
        Inches(x * XI), Inches(y * YI), Inches(w * XI), Inches(h * YI))
    if round_:
        try:
            shp.adjustments[0] = round_
        except Exception:
            pass
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = hx(fill)
        if alpha_pct:
            alpha(shp, alpha_pct)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = hx(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    shp.text_frame.word_wrap = False
    return shp


def ellipse(slide, x, y, w, h, fill=None, line=None, lw=1.0, alpha_pct=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x * XI), Inches(y * YI),
                                 Inches(w * XI), Inches(h * YI))
    if fill:
        shp.fill.solid()
        shp.fill.fore_color.rgb = hx(fill)
        if alpha_pct:
            alpha(shp, alpha_pct)
    else:
        shp.fill.background()
    if line:
        shp.line.color.rgb = hx(line)
        shp.line.width = Pt(lw)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def leaf(slide, x, y, w, h, rot, color, a=0.45):
    shp = slide.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(x * XI), Inches(y * YI),
                                 Inches(w * XI), Inches(h * YI))
    shp.rotation = rot
    shp.fill.solid()
    shp.fill.fore_color.rgb = hx(color)
    alpha(shp, a)
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def set_font(run, name=FONT):
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)


def text(slide, lines, x, y, w, h=None, size=18, color=C_TEXT, bold=False,
         align=PP_ALIGN.LEFT, line_pt=None, anchor=MSO_ANCHOR.TOP):
    """lines: list of str or list of (str, color)"""
    tb = slide.shapes.add_textbox(Inches(x * XI), Inches(y * YI),
                                  Inches(w * XI), Inches((h or size / PT) * YI))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = Pt(line_pt) if line_pt else 1.0
    for i, ln in enumerate(lines):
        col = color
        if isinstance(ln, tuple):
            ln, col = ln
        if i:
            p._p.append(p._p.makeelement(qn("a:br"), {}))
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.color.rgb = hx(col)
        set_font(r)
    return tb


def card(slide, x, y, w, h, txt, size=30, color=C_WHITE, fill=C_GREEN, round_=0.08):
    rect(slide, x, y, w, h, fill=fill, round_=round_)
    text(slide, [txt], x, y, w, h, size=size, color=color, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)


def header(slide, title):
    text(slide, [title], 58, 61, 400, 60, size=34, color=C_TITLE, bold=True)


def decor(slide):
    """水墨风装饰：左下橙水彩、右上淡绿晕染、右侧竹叶（全部为可编辑形状）"""
    ellipse(slide, -60, 470, 300, 230, fill="E9A97E", alpha_pct=0.62)
    ellipse(slide, 20, 505, 190, 170, fill="D9764A", alpha_pct=0.72)
    ellipse(slide, 900, -70, 300, 240, fill="DDE3D4", alpha_pct=0.55)
    ellipse(slide, 980, 120, 190, 150, fill="C6D3BE", alpha_pct=0.7)
    for dx, dy, rw, rh, rot in ((905, 40, 150, 210, 25), (955, 130, 130, 190, -35),
                                (855, 130, 120, 170, 55)):
        leaf(slide, dx, dy, rw, rh, rot, "93A57C", a=0.55)


def asset(page, box, name):
    """crop a photo out of the original frame so it becomes a replaceable picture"""
    img = cv2.imread(os.path.join(W, "pages_bg", f"p{page:02d}_raw.png"))
    img = cv2.resize(img, (1106, 648))
    x, y, w, h = box
    crop = img[max(0, y):y + h, max(0, x):x + w]
    p = os.path.join(ASSETS, name)
    cv2.imwrite(p, crop, [cv2.IMWRITE_PNG_COMPRESSION, 3])
    return p


def pic(slide, path, x, y, w, h):
    return slide.shapes.add_picture(path, Inches(x * XI), Inches(y * YI),
                                    Inches(w * XI), Inches(h * YI))


def newpage(title=None):
    s = prs.slides.add_slide(BLANK)
    rect(s, 0, 0, 1106, 648, fill=C_BG)
    decor(s)
    if title:
        header(s, title)
    return s


PT_PX = PT   # px -> pt factor

# =============================== 1 封面 ===============================
s = newpage()
text(s, ["人教版  八年级物理  第四章第二节"], 39, 87, 620, 40, size=25, color="6D7853")
text(s, ["《光的反射》说课"], 150, 278, 806, 110, size=71, color=C_ORANGE, bold=True,
     align=PP_ALIGN.CENTER)

# =============================== 2 目录 ===============================
s = newpage()
text(s, ["目录"], 128, 70, 260, 60, size=45, color=C_TITLE, bold=True)
for t, x, y in [("1.教材分析", 190, 191), ("2.学情分析", 189, 318), ("3.教学目标", 191, 446),
                ("4.教学方法", 575, 193), ("5.教学过程", 576, 319), ("6.板书设计", 575, 445)]:
    text(s, [t], x, y, 340, 60, size=41, color=C_OCHRE, bold=True)

# =============================== 3 教材分析·教材图片 ===============================
s = newpage("教材分析")
text(s, ["人教版初中物理八年级上册"], 597, 115, 420, 34, size=24, color=C_TEXT)
text(s, ["第四章"], 212, 174, 260, 50, size=37, color=C_TEXT, bold=True)
text(s, ["第一节光的直线传播",
         ("第二节 光的反射", C_RED),
         "第三节平面镜成像",
         "第四节光的折射",
         "第五节光的色散"],
     95, 243, 420, 300, size=31, color=C_TEXT, line_pt=(498 - 243) / 4 * PT_PX)
pic(s, asset(3, (631, 178, 313, 415), "p03_book.png"), 631, 178, 313, 415)

# =============================== 4 教材分析·作用和地位 ===============================
s = newpage("教材分析")
text(s, ["人教版初中物理八年级上册"], 608, 86, 420, 32, size=23, color=C_TEXT)
text(s, ["第四章"], 174, 136, 260, 50, size=37, color=C_TEXT, bold=True)
text(s, ["第一节光的直线传播",
         ("第二节 光的反射", C_RED),
         "第三节平面镜成像",
         "第四节光的折射",
         "第五节光的色散"],
     55, 204, 420, 300, size=31, color=C_TEXT, line_pt=(462 - 204) / 4 * PT_PX)
rect(s, 595, 157, 406, 361, fill=C_GREEN, round_=0.06)
text(s, ["作用和地位"], 660, 200, 280, 50, size=29, color=C_WHITE, bold=True, align=PP_ALIGN.CENTER)
text(s, ["1.内容生动有趣，贴近我们的", "日常生活",
         "2.继光的直线传播后对光现象", "的进一步完善",
         "3.为后面平面镜成像的学习打", "下坚实的知识基础"],
     618, 272, 372, 190, size=21, color=C_WHITE, line_pt=(451 - 278) / 5 * PT_PX)
text(s, ["教学重点：通过实验探究了解光的反射规律。"], 54, 553, 950, 48, size=34, color=C_RED, bold=True)

# =============================== 5 学情分析 ===============================
s = newpage("学情分析")
for t, x in [("知识基础", 181), ("心理特点", 509), ("学习困难", 814)]:
    text(s, [t], x - 10, 190, 160, 34, size=21, color=C_TEXT, bold=True)
text(s, ["1.已经学习了光的直线", "传播", "2.对光的反射也有了初", "步的了解"],
     82, 260, 285, 170, size=19, color=C_TEXT, line_pt=(390 - 260) / 3 * PT_PX)
text(s, ["1.对有趣的物理现象充", "满了好奇心", "2.具有强烈的求知欲望"],
     417, 261, 285, 140, size=19, color=C_TEXT, line_pt=(348 - 261) / 2 * PT_PX)
text(s, ["1.空间思维能力有限", "2.对抽象概念的理解存", "在困难"],
     753, 260, 285, 140, size=19, color=C_TEXT, line_pt=(345 - 260) / 2 * PT_PX)
text(s, ["教学难点：掌握光的反射规律，并运用该规律", "来解释生活中的现象"],
     69, 468, 800, 110, size=34, color=C_RED, bold=True, line_pt=50 * PT_PX)

# =============================== 6 教学目标 ===============================
s = newpage("教学目标")
ellipse(s, 390, 225, 298, 298, line=C_GREENL, lw=26)
ellipse(s, 470, 305, 138, 138, line="C6D3BE", lw=6)
for t, x, y in [("1.物理观念", 226, 193), ("3.科学探究", 717, 192),
                ("2.科学思维", 224, 418), ("4.科学态度与责任", 719, 419)]:
    text(s, [t], x - 6, y, 240, 36, size=24, color=C_OCHRE, bold=True)
text(s, ["学生应理解光的反射概念；", "认识光反射的规律；能够用", "光的反射定律解释生活现象。"],
     89, 245, 300, 110, size=18, color=C_TEXT, line_pt=36 * PT_PX)
text(s, ["通过动手操作和实验探究，", "使学生经历基本的科学探究", "过程，学习探究科学的方法。"],
     716, 244, 370, 110, size=18, color=C_TEXT, line_pt=36 * PT_PX)
text(s, ["学生会运用课堂所学知识来", "解释生活中有关反射的现象。"],
     89, 472, 300, 75, size=18, color=C_TEXT, line_pt=36 * PT_PX)
text(s, ["培养学生的动手能力及分析能力；", "感悟科学、技术和社会的相互联系。"],
     717, 471, 370, 75, size=18, color=C_TEXT, line_pt=37 * PT_PX)

# =============================== 7 教学方法 ===============================
s = newpage("教学方法")
text(s, ["以教师为主导，学生为主体"], 301, 76, 460, 40, size=27, color="813815", bold=True,
     align=PP_ALIGN.CENTER)
text(s, ["教法"], 260, 175, 140, 55, size=37, color="833611", bold=True, align=PP_ALIGN.CENTER)
text(s, ["学法"], 660, 179, 140, 55, size=37, color="833611", bold=True, align=PP_ALIGN.CENTER)
for y in (312, 405, 498):
    card(s, 178, y, 296, 94, "", fill=C_GREEN)
    card(s, 578, y, 297, 94, "", fill=C_GREEN)
for t, x, y in [("情景激学", 178, 312), ("启发教学", 178, 405), ("讲授分析", 178, 498),
                ("实验探究", 578, 312), ("对比归纳", 578, 405), ("分析总结", 578, 498)]:
    text(s, [t], x, y, 296, 94, size=30, color=C_WHITE, bold=True,
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# =============================== 8 教学资源 ===============================
s = newpage("教学方法")
for cx, cy in ((553, 217), (682, 291), (682, 440), (425, 291), (425, 440), (553, 514)):
    ellipse(s, cx - 48, cy - 48, 96, 96, fill=C_GREEN)
text(s, ["资源"], 506, 324, 130, 60, size=44, color=C_GREEN, bold=True, align=PP_ALIGN.CENTER)
for t, x, y, w in [("实验室常见仪器", 147, 205, 210), ("视频", 768, 203, 90),
                   ("生活废弃物", 202, 333, 180), ("图片", 769, 330, 90),
                   ("自制实验装置", 176, 462, 190), ("多媒体课件", 772, 461, 160)]:
    text(s, [t], x, y, w, 36, size=23, color=C_TEXT)

# =============================== 9 教学过程·总览 ===============================
s = newpage("教学过程")
ellipse(s, 337, 180, 433, 365, line=C_GREENL, lw=22)
ellipse(s, 430, 272, 247, 208, line="CAD5B9", lw=5)
text(s, ["教学过程"], 470, 500, 170, 40, size=24, color=C_GREEN, bold=True, align=PP_ALIGN.CENTER)
for t, x, y in [("1.创设情境，引入新课", 119, 180), ("2.知识回顾，引出概念", 69, 322),
                ("3.实验探究，得出结论", 113, 472), ("4.课堂揭秘，巩固提升", 769, 182),
                ("5.联系生活，学以致用", 813, 322), ("6.课外拓展，感知生活", 769, 473)]:
    text(s, [t], x, y, 300, 34, size=20, color=C_TEXT)

# =============================== 10 环节一 ===============================
s = newpage("教学过程")
text(s, ["1.创设情境，引入新课"], 267, 71, 330, 40, size=28, color=C_TEXT, bold=True)
pic(s, asset(10, (100, 180, 451, 304), "p10_box.png"), 100, 180, 451, 304)
text(s, ["魔术小盒"], 250, 500, 160, 34, size=21, color=C_TEXT)
rect(s, 672, 184, 322, 296, fill=C_GREEN, round_=0.06)
text(s, ["魔术小盒", "中究竟藏", "着什么秘", "密呢?"], 700, 205, 270, 250, size=34,
     color=C_WHITE, bold=True, align=PP_ALIGN.CENTER, line_pt=56 * PT_PX)

# =============================== 11 环节二 ===============================
s = newpage("教学过程")
text(s, ["2.知识回顾，引出概念"], 266, 70, 330, 40, size=28, color=C_TEXT, bold=True)
pic(s, asset(11, (107, 152, 431, 336), "p11_ray.png"), 107, 152, 431, 336)
text(s, ["反射现象"], 250, 500, 160, 34, size=21, color=C_TEXT)
rect(s, 665, 154, 323, 311, fill=C_GREEN, round_=0.06)
text(s, ["知识回顾，", "观察实验现", "象，引出反", "射的概念。"], 690, 190, 280, 250, size=30,
     color=C_WHITE, bold=True, align=PP_ALIGN.CENTER, line_pt=54 * PT_PX)
text(s, ["思考：为什么不发光的物体也能被我们看见呢?"], 101, 556, 930, 44, size=30, color="793413",
     bold=True, align=PP_ALIGN.CENTER)

# =============================== 12 环节三 ===============================
s = newpage("教学过程")
text(s, ["3.实验探究，得出结论"], 267, 72, 330, 40, size=28, color=C_TEXT, bold=True)
pic(s, asset(12, (143, 146, 348, 342), "p12_device.png"), 143, 146, 348, 342)
text(s, ["反射演示仪"], 240, 498, 160, 34, size=21, color=C_TEXT)
rect(s, 596, 160, 428, 318, fill=C_GREEN, round_=0.06)
text(s, ["得出准确结论"], 620, 185, 380, 50, size=35, color=C_WHITE, bold=True, align=PP_ALIGN.CENTER)
text(s, ["1.法线居中", "2.三线共面", "3.两角相等", "4.光路可逆"], 640, 240, 340, 210, size=30,
     color=C_WHITE, bold=True, align=PP_ALIGN.CENTER, line_pt=44 * PT_PX)
text(s, ["实验：分组实验，探究光的反射规律有哪些?"], 103, 556, 930, 44, size=30, color="833B1C",
     bold=True, align=PP_ALIGN.CENTER)

# =============================== 13 环节四 ===============================
s = newpage("教学过程")
text(s, ["4.课堂揭秘，巩固提升"], 268, 72, 330, 40, size=28, color=C_TEXT, bold=True)
pic(s, asset(13, (96, 198, 432, 280), "p13_left.png"), 96, 198, 432, 280)
text(s, ["魔术小盒"], 250, 510, 160, 34, size=21, color=C_TEXT)
rect(s, 582, 206, 404, 262, fill=C_CYAN, round_=0.06)
text(s, ["原理解释"], 730, 510, 160, 34, size=21, color=C_TEXT)

# =============================== 14 环节五 ===============================
s = newpage("教学过程")
text(s, ["5.联系生活，学以致用"], 268, 70, 330, 40, size=28, color=C_TEXT, bold=True)
pic(s, asset(14, (45, 194, 345, 277), "p14_a.png"), 45, 194, 345, 277)
pic(s, asset(14, (427, 194, 257, 277), "p14_b.png"), 427, 194, 257, 277)
pic(s, asset(14, (728, 194, 348, 277), "p14_c.png"), 728, 194, 348, 277)
for t, x in [("汽车反光镜", 163), ("潜望镜", 521), ("自制潜望镜", 837)]:
    text(s, [t], x - 15, 493, 170, 34, size=21, color=C_TEXT, align=PP_ALIGN.CENTER)

# =============================== 15 环节六 ===============================
s = newpage("教学过程")
text(s, ["6.课外拓展，感知生活"], 267, 72, 330, 40, size=28, color=C_TEXT, bold=True)
pic(s, asset(15, (238, 190, 547, 300), "p15_video.png"), 238, 190, 547, 300)
text(s, ["“隐身术”视频"], 410, 512, 220, 34, size=21, color=C_TEXT, align=PP_ALIGN.CENTER)

# =============================== 16 板书设计 ===============================
s = newpage("板书设计")
rect(s, 160, 172, 792, 378, fill=C_BOARD, round_=0.05)
text(s, ["光的反射"], 420, 200, 280, 60, size=39, color=C_WHITE, bold=True, align=PP_ALIGN.CENTER)
text(s, ["1.", "2.", "3."], 355, 310, 82, 160, size=32, color=C_WHITE, bold=True,
     align=PP_ALIGN.RIGHT, line_pt=54 * PT_PX)
text(s, ["光的反射现象", "光的反射规律", "光的反射应用"], 450, 310, 300, 160, size=39,
     color=C_WHITE, bold=True, line_pt=54 * PT_PX)

# =============================== 17 致谢 ===============================
s = newpage()
text(s, ["感谢您的耐心聆听"], 150, 284, 806, 80, size=54, color=C_ORANGE, bold=True,
     align=PP_ALIGN.CENTER)

prs.save(OUT)
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
